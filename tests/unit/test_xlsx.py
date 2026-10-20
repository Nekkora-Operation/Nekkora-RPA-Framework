"""0.3.0 — escrita genérica de XLSX: esquema, saneamento, anti-fórmula, gravação atômica, destino transacional e
consolidação pelo engine. A conferência é sempre feita REABRINDO o arquivo (openpyxl e o zip cru)."""

from __future__ import annotations

import json
import os
import zipfile
from pathlib import Path

import openpyxl
import pytest
from tests.unit.test_engine import config

import nekkora_rpa.xlsx as xlsx_mod
from nekkora_rpa import (
    Engine,
    InMemoryInputProvider,
    MemoryEventSink,
    OutputRun,
    OutputStatus,
    ProcessResult,
    RpaProcess,
    RunStatus,
    Runtime,
    TransactionItem,
    TransactionStatus,
)
from nekkora_rpa.core.contracts import SinkRejected, TransactionOutcome
from nekkora_rpa.errors import BusinessRuleRejected, TransientTechnicalFailure
from nekkora_rpa.xlsx import (
    ColumnKind,
    CommitPolicy,
    OutputColumn,
    OverwritePolicy,
    XlsxOutputError,
    XlsxResultSink,
    XlsxSchema,
    prepare_row,
    read_back,
    write_xlsx,
)

SCHEMA = XlsxSchema(
    (
        OutputColumn("URL", "url", ColumnKind.URL),
        OutputColumn("Nome completo", "nome"),
        OutputColumn("Nascimento", "nascimento", ColumnKind.ISO_DATE),
        OutputColumn("Morte", "morte", ColumnKind.ISO_DATE, required=False),
        OutputColumn("Ocupação", "ocupacao", max_chars=50),
    ),
    key_field="url",
    sheet="Biografias",
)
SENT = "SENTINELA-CELULA-7f3a"


def row(n: int, **extra: object) -> dict[str, object]:
    base: dict[str, object] = {
        "url": f"https://exemplo.test/p/{n}",
        "nome": f"Pessoa {n} Ünicode ção 日本",
        "nascimento": "1839-06-21",
        "morte": "",
        "ocupacao": "escritora; jornalista",
    }
    base.update(extra)
    return base


def raw_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".xml"))


# --------------------------------------------------------------------------------------------- esquema
@pytest.mark.parametrize(
    ("columns", "key", "code"),
    [
        ((), "url", "schema_invalid"),
        ((OutputColumn("URL", "url"), OutputColumn("url", "outro")), "url", "schema_header_invalid"),  # sem caixa
        ((OutputColumn("=URL", "url"),), "url", "schema_header_invalid"),
        ((OutputColumn(" URL", "url"),), "url", "schema_header_invalid"),
        ((OutputColumn("URL", "URL"),), "URL", "schema_field_invalid"),
        ((OutputColumn("URL", "url"), OutputColumn("B", "url")), "url", "schema_field_invalid"),
        ((OutputColumn("URL", "url", required=False),), "url", "schema_key_invalid"),
        ((OutputColumn("URL", "url"),), "outro", "schema_key_invalid"),
        ((OutputColumn("URL", "url", max_chars=0),), "url", "schema_limit_invalid"),
        ((OutputColumn("URL", "url", max_chars=40_000),), "url", "schema_limit_invalid"),
    ],
)
def test_esquema_invalido_e_recusado(columns: tuple[OutputColumn, ...], key: str, code: str) -> None:
    with pytest.raises(XlsxOutputError) as e:
        XlsxSchema(columns, key_field=key)
    assert e.value.code == code


@pytest.mark.parametrize("sheet", ["", "a" * 32, "x/y", "x[1]", " Dados", "'Dados"])
def test_nome_de_aba_invalido(sheet: str) -> None:
    with pytest.raises(XlsxOutputError, match="schema_sheet_invalid"):
        XlsxSchema((OutputColumn("URL", "url"),), key_field="url", sheet=sheet)


def test_revisao_do_esquema_muda_com_qualquer_coluna() -> None:
    other = XlsxSchema(
        (*SCHEMA.columns[:-1], OutputColumn("Ocupação", "ocupacao", max_chars=51)), key_field="url", sheet="Biografias"
    )
    assert SCHEMA.revision() != other.revision()
    assert SCHEMA.revision() == XlsxSchema(SCHEMA.columns, key_field="url", sheet="Biografias").revision()


# --------------------------------------------------------------------------------------------- linhas
def test_linha_saneada_na_ordem_das_colunas() -> None:
    r = prepare_row(SCHEMA, row(1, nome="  Ana \x07Maria\x00  Café ", ocupacao=" a\n b "))
    assert r == ("https://exemplo.test/p/1", "Ana Maria Café", "1839-06-21", "", "a b")


@pytest.mark.parametrize(
    ("data", "code"),
    [
        ({"nome": "x", "nascimento": "1839-06-21", "ocupacao": "x"}, "missing_url"),
        (row(1, nome="  "), "missing_nome"),
        (row(1, nascimento="21/06/1839"), "date_invalid_nascimento"),
        (row(1, nascimento="1900-02-30"), "date_invalid_nascimento"),
        (row(1, morte="junho de 1839"), "date_invalid_morte"),
        (row(1, url="javascript:alert(1)"), "url_invalid_url"),
        (row(1, ocupacao="x" * 51), "too_long_ocupacao"),
        (row(1, nome=12), "value_type_nome"),
        (row(1, extra="x"), "unknown_field"),
    ],
)
def test_linha_incompleta_ou_invalida(data: dict[str, object], code: str) -> None:
    with pytest.raises(XlsxOutputError) as e:
        prepare_row(SCHEMA, data)
    assert e.value.code == code
    assert SENT not in str(e.value)


# --------------------------------------------------------------------------------------------- arquivo
def test_escrita_reaberta_unicode_datas_como_texto(tmp_path: Path) -> None:
    out = tmp_path / "saida.xlsx"
    rep = write_xlsx(out, SCHEMA, [row(1), row(2, morte="1908-09-29")])
    assert rep.status is OutputStatus.COMMITTED and rep.rows == 2
    assert read_back(out, SCHEMA) == [
        ("https://exemplo.test/p/1", "Pessoa 1 Ünicode ção 日本", "1839-06-21", "", "escritora; jornalista"),
        ("https://exemplo.test/p/2", "Pessoa 2 Ünicode ção 日本", "1839-06-21", "1908-09-29", "escritora; jornalista"),
    ]
    wb = openpyxl.load_workbook(out)
    ws = wb["Biografias"]
    assert wb.sheetnames == ["Biografias"]
    assert [c.value for c in ws[1]] == list(SCHEMA.headers)
    assert all(c.data_type == "s" for r in ws.iter_rows(min_row=2) for c in r if c.value is not None)
    assert ws["C2"].number_format == "@"  # data gravada como TEXTO (1839 nem existe no calendário do Excel)
    assert ws.max_row == 3 and ws.max_column == 5
    assert not list(tmp_path.glob(".*tmp*"))  # nenhum temporário sobra


def test_formula_injection_vira_texto(tmp_path: Path) -> None:
    out = tmp_path / "f.xlsx"
    perigosos = ['=HYPERLINK("http://x.invalid","y")', "+1+1", "-2+3", "@SUM(A1)", "=1+1"]
    write_xlsx(out, SCHEMA, [row(i, nome=v) for i, v in enumerate(perigosos)])
    xml = raw_xml(out)
    assert "<f>" not in xml and "<f " not in xml  # nenhuma fórmula no pacote
    ws = openpyxl.load_workbook(out)["Biografias"]
    assert [ws.cell(row=i + 2, column=2).value for i in range(len(perigosos))] == perigosos
    assert all(ws.cell(row=i + 2, column=2).quotePrefix for i in range(len(perigosos)))
    assert [r[1] for r in read_back(out, SCHEMA)] == perigosos


def test_duplicidade_e_excesso_recusados_antes_do_disco(tmp_path: Path) -> None:
    out = tmp_path / "d.xlsx"
    with pytest.raises(XlsxOutputError, match="duplicate_key"):
        write_xlsx(out, SCHEMA, [row(1), row(1, nome="outro")])
    small = XlsxSchema(SCHEMA.columns, key_field="url", sheet="Biografias", max_rows=1)
    with pytest.raises(XlsxOutputError, match="too_many_rows"):
        write_xlsx(out, small, [row(1), row(2)])
    assert not out.exists()


def test_politica_de_sobrescrita(tmp_path: Path) -> None:
    out = tmp_path / "s.xlsx"
    write_xlsx(out, SCHEMA, [row(1)])
    before = out.read_bytes()
    with pytest.raises(XlsxOutputError, match="exists"):
        write_xlsx(out, SCHEMA, [row(2)])
    assert out.read_bytes() == before
    write_xlsx(out, SCHEMA, [row(2)], overwrite=OverwritePolicy.REPLACE)
    assert read_back(out, SCHEMA)[0][0] == "https://exemplo.test/p/2"


@pytest.mark.parametrize(("name", "code"), [("x.csv", "path_not_xlsx"), ("falta/x.xlsx", "folder_missing")])
def test_caminho_invalido(tmp_path: Path, name: str, code: str) -> None:
    with pytest.raises(XlsxOutputError, match=code):
        write_xlsx(tmp_path / name, SCHEMA, [row(1)])


def test_arquivo_travado_mantem_o_anterior_e_nao_deixa_temporario(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    out = tmp_path / "t.xlsx"
    write_xlsx(out, SCHEMA, [row(1)])
    before = out.read_bytes()

    def locked(src: object, dst: object) -> None:
        raise PermissionError(13, "em uso")  # o que o Windows devolve com o arquivo aberto no Excel

    monkeypatch.setattr(xlsx_mod, "_replace", locked)
    with pytest.raises(XlsxOutputError, match="locked"):
        write_xlsx(out, SCHEMA, [row(9)], overwrite=OverwritePolicy.REPLACE)
    assert out.read_bytes() == before
    assert not [p for p in tmp_path.iterdir() if p.name.startswith(".")]


@pytest.mark.skipif(os.name != "nt", reason="trava de arquivo real do Windows")
def test_arquivo_travado_de_verdade_no_windows(tmp_path: Path) -> None:
    import msvcrt

    out = tmp_path / "w.xlsx"
    write_xlsx(out, SCHEMA, [row(1)])
    before = out.read_bytes()
    with out.open("r+b") as fh:  # trava exclusiva como a de um editor
        msvcrt.locking(fh.fileno(), msvcrt.LK_NBLCK, 1)
        try:
            with pytest.raises(XlsxOutputError, match=r"locked|io_error"):
                write_xlsx(out, SCHEMA, [row(2)], overwrite=OverwritePolicy.REPLACE)
        finally:
            fh.seek(0)
            msvcrt.locking(fh.fileno(), msvcrt.LK_UNLCK, 1)
    assert out.read_bytes() == before


def test_conferencia_divergente_nao_publica(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import nekkora_rpa.xlsx as mod

    out = tmp_path / "v.xlsx"
    monkeypatch.setattr(mod, "read_back", lambda p, s: [("adulterado",)])
    with pytest.raises(XlsxOutputError, match="verify_mismatch"):
        write_xlsx(out, SCHEMA, [row(1)])
    assert not out.exists()


def test_leitura_recusa_formula_aba_extra_e_cabecalho_diferente(tmp_path: Path) -> None:
    p = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Biografias"
    ws.append(list(SCHEMA.headers))
    ws.append(["https://e.test/1", "=1+1", "1839-06-21", None, "x"])
    wb.save(p)
    with pytest.raises(XlsxOutputError, match="verify_formula"):
        read_back(p, SCHEMA)
    wb.create_sheet("Outra")
    wb["Biografias"]["B2"] = "texto"
    wb.save(p)
    with pytest.raises(XlsxOutputError, match="verify_sheet"):
        read_back(p, SCHEMA)
    del wb["Outra"]
    wb["Biografias"]["A1"] = "Endereço"
    wb.save(p)
    with pytest.raises(XlsxOutputError, match="verify_header"):
        read_back(p, SCHEMA)


# --------------------------------------------------------------------------------------------- destino transacional
def outcome(n: int, status: TransactionStatus = TransactionStatus.SUCCEEDED, **extra: object) -> TransactionOutcome:
    return TransactionOutcome(
        f"tx-{n:020d}", status, row(n, **extra) if status is TransactionStatus.SUCCEEDED else None
    )


ITEM = TransactionItem("ref", {})


def sink(tmp_path: Path, **kw: object) -> XlsxResultSink:
    return XlsxResultSink(tmp_path / "out" / "bio.xlsx", SCHEMA, state_dir=tmp_path / "estado", **kw)  # type: ignore[arg-type]


def test_tudo_ou_nada_retem_planilha_parcial(tmp_path: Path) -> None:
    (tmp_path / "out").mkdir()
    s = sink(tmp_path)
    s.deliver(ITEM, outcome(1))
    s.deliver(ITEM, outcome(2, TransactionStatus.FAILED_TECHNICAL))
    rep = s.finalize(OutputRun(True, ((outcome(1).transaction_id, TransactionStatus.SUCCEEDED),)))
    assert rep.status is OutputStatus.WITHHELD and rep.detail == "transactions_not_succeeded"
    assert not (tmp_path / "out" / "bio.xlsx").exists()
    assert s.state == "WITHHELD"
    rep = sink(tmp_path).finalize(OutputRun(False))
    assert rep.status is OutputStatus.WITHHELD


def test_somente_sucesso_grava_as_boas(tmp_path: Path) -> None:
    (tmp_path / "out").mkdir()
    s = sink(tmp_path, commit=CommitPolicy.SUCCEEDED_ONLY)
    s.deliver(ITEM, outcome(1))
    s.deliver(ITEM, outcome(2, TransactionStatus.BUSINESS_REJECTED))
    assert s.finalize(OutputRun(True)).status is OutputStatus.COMMITTED
    assert len(read_back(tmp_path / "out" / "bio.xlsx", SCHEMA)) == 1


def test_linha_ruim_e_chave_repetida_sao_recusadas_sem_preparar(tmp_path: Path) -> None:
    s = sink(tmp_path)
    with pytest.raises(SinkRejected, match="missing_nome"):
        s.deliver(ITEM, outcome(1, nome=""))
    s.deliver(ITEM, outcome(1))
    dup = TransactionOutcome("tx-" + "9" * 20, TransactionStatus.SUCCEEDED, row(1))
    with pytest.raises(SinkRejected, match="duplicate_key"):
        s.deliver(ITEM, dup)
    assert s.staged_rows == 1
    s.deliver(ITEM, outcome(1))  # mesma transação de novo: idempotente
    assert s.staged_rows == 1


def test_quantidade_esperada(tmp_path: Path) -> None:
    (tmp_path / "out").mkdir()
    s = sink(tmp_path, expected_rows=2)
    s.deliver(ITEM, outcome(1))
    assert s.finalize(OutputRun(True)).detail == "row_count_mismatch"
    s.deliver(ITEM, outcome(2))
    assert s.finalize(OutputRun(True)).status is OutputStatus.COMMITTED


def test_queda_antes_do_commit_retoma_sem_duplicar(tmp_path: Path) -> None:
    (tmp_path / "out").mkdir()
    s = sink(tmp_path)
    s.deliver(ITEM, outcome(1))
    s.deliver(ITEM, outcome(2))
    del s  # "queda": nada foi consolidado
    s2 = sink(tmp_path)
    assert s2.staged_rows == 2 and s2.state == "OPEN"
    s2.deliver(ITEM, outcome(2))  # a retomada entrega de novo uma transação já preparada
    s2.deliver(ITEM, outcome(3))
    assert s2.finalize(OutputRun(True)).status is OutputStatus.COMMITTED
    assert [r[0] for r in read_back(tmp_path / "out" / "bio.xlsx", SCHEMA)] == [
        f"https://exemplo.test/p/{n}" for n in (1, 2, 3)
    ]


def test_queda_entre_conferencia_e_troca_e_reconciliada(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    (tmp_path / "out").mkdir()
    s = sink(tmp_path)
    s.deliver(ITEM, outcome(1))

    class Queda(BaseException):
        pass

    real = xlsx_mod._replace

    def cai(src: object, dst: object) -> None:
        raise Queda()

    monkeypatch.setattr(xlsx_mod, "_replace", cai)
    with pytest.raises(Queda):
        s.finalize(OutputRun(True))  # "processo morto" logo antes da troca
    monkeypatch.setattr(xlsx_mod, "_replace", real)
    state = json.loads((tmp_path / "estado" / "bio.xlsx-output.json").read_text(encoding="utf-8"))
    assert state["state"] == "COMMITTING" and state["pending_sha256"]
    assert not (tmp_path / "out" / "bio.xlsx").exists()
    s2 = sink(tmp_path)  # retomada: arquivo final ausente → volta a OPEN e grava
    assert s2.state == "OPEN"
    assert s2.finalize(OutputRun(True)).status is OutputStatus.COMMITTED


def test_queda_depois_da_troca_e_reconhecida_pelo_hash(tmp_path: Path) -> None:
    (tmp_path / "out").mkdir()
    s = sink(tmp_path)
    s.deliver(ITEM, outcome(1))
    assert s.finalize(OutputRun(True)).status is OutputStatus.COMMITTED
    st = tmp_path / "estado" / "bio.xlsx-output.json"
    doc = json.loads(st.read_text(encoding="utf-8"))
    doc["state"], doc["pending_sha256"], doc["committed_sha256"] = "COMMITTING", doc["committed_sha256"], None
    st.write_text(json.dumps(doc), encoding="utf-8")  # trocou o arquivo e caiu antes de registrar COMMITTED
    before = (tmp_path / "out" / "bio.xlsx").read_bytes()
    s2 = sink(tmp_path)
    assert s2.state == "COMMITTED"
    rep = s2.finalize(OutputRun(True))
    assert rep.status is OutputStatus.COMMITTED and rep.detail == "already_committed"
    assert (tmp_path / "out" / "bio.xlsx").read_bytes() == before  # não regravou


def test_estado_de_outro_esquema_nunca_mistura(tmp_path: Path) -> None:
    s = sink(tmp_path)
    s.deliver(ITEM, outcome(1))
    other = XlsxSchema(SCHEMA.columns[:3], key_field="url", sheet="Biografias")
    with pytest.raises(XlsxOutputError, match="state_mismatch"):
        XlsxResultSink(tmp_path / "out" / "bio.xlsx", other, state_dir=tmp_path / "estado")


def test_valores_nao_aparecem_em_erros_nem_no_resumo(tmp_path: Path) -> None:
    s = sink(tmp_path)
    with pytest.raises(SinkRejected) as e:
        s.deliver(ITEM, outcome(1, ocupacao=SENT * 10))
    assert SENT not in str(e.value)


# --------------------------------------------------------------------------------------------- pelo engine
class Bio(RpaProcess):
    name = "bio"
    requires_browser = False

    def __init__(self, fail: set[str] | None = None) -> None:
        self.fail = fail or set()

    def process(self, rt: Runtime, item: TransactionItem) -> ProcessResult:
        n = int(item.reference)
        if item.reference in self.fail:
            raise TransientTechnicalFailure(detail="indisponivel")
        if n == 99:
            raise BusinessRuleRejected("SEM_DADOS")
        return ProcessResult(row(n, nome=f"{SENT} {n}"))


def engine_run(
    tmp_path: Path, refs: list[str], proc: Bio, **kw: object
) -> tuple[object, MemoryEventSink, XlsxResultSink]:
    (tmp_path / "out").mkdir(exist_ok=True)
    events = MemoryEventSink()
    s = sink(tmp_path, expected_rows=len(refs), **kw)
    from nekkora_rpa import CheckpointStore

    eng = Engine(
        proc,
        InMemoryInputProvider([TransactionItem(r, {}) for r in refs]),
        s,
        config(),
        run_dir=tmp_path / "run",
        event_sinks=[events],
        checkpoint=CheckpointStore(tmp_path / "run" / "checkpoint.json", "bio"),
    )
    return eng.run(), events, s


def test_engine_consolida_uma_vez_depois_do_lote(tmp_path: Path) -> None:
    summary, events, _ = engine_run(tmp_path, ["1", "2", "3"], Bio())
    assert summary.status is RunStatus.COMPLETED  # type: ignore[attr-defined]
    assert summary.output.status is OutputStatus.COMMITTED and summary.output.rows == 3  # type: ignore[attr-defined]
    codes = events.codes()
    assert codes.count("OUTPUT_COMMITTED") == 1
    assert codes.index("OUTPUT_COMMITTED") > max(i for i, c in enumerate(codes) if c == "TRANSACTION_COMPLETED")
    rows = read_back(tmp_path / "out" / "bio.xlsx", SCHEMA)
    assert [r[0] for r in rows] == [f"https://exemplo.test/p/{n}" for n in (1, 2, 3)]
    summ = json.loads((tmp_path / "run" / "summary.json").read_text(encoding="utf-8"))
    assert summ["version"] == "1.2.0" and summ["output"] == {"status": "COMMITTED", "rows": 3, "detail": None}
    dump = (tmp_path / "run" / "summary.json").read_text(encoding="utf-8") + json.dumps(
        [e.to_dict() for e in events.events]
    )
    assert SENT not in dump  # nenhuma célula em eventos nem no resumo


def test_engine_com_falha_retem_e_a_retomada_completa_sem_duplicar(tmp_path: Path) -> None:
    summary, events, _ = engine_run(tmp_path, ["1", "2", "3"], Bio(fail={"2"}))
    assert summary.status is RunStatus.COMPLETED_WITH_FAILURES  # type: ignore[attr-defined]
    assert summary.output.status is OutputStatus.WITHHELD  # type: ignore[attr-defined]
    assert "OUTPUT_WITHHELD" in events.codes()
    assert not (tmp_path / "out" / "bio.xlsx").exists()  # nenhuma planilha parcial
    summary2, _, s2 = engine_run(tmp_path, ["1", "2", "3"], Bio())  # retomada: 1 e 3 já finais, 2 reprocessada
    assert summary2.output.status is OutputStatus.COMMITTED  # type: ignore[attr-defined]
    rows = read_back(tmp_path / "out" / "bio.xlsx", SCHEMA)
    assert len(rows) == 3 and len({r[0] for r in rows}) == 3
    assert s2.state == "COMMITTED"


def test_engine_com_rejeicao_de_negocio_nao_grava(tmp_path: Path) -> None:
    summary, _, _ = engine_run(tmp_path, ["1", "99"], Bio())
    assert summary.output.status is OutputStatus.WITHHELD  # type: ignore[attr-defined]
    assert summary.status is RunStatus.COMPLETED_WITH_FAILURES  # type: ignore[attr-defined]


def test_engine_com_saida_travada_termina_com_falha_registrada(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    def locked(src: object, dst: object) -> None:
        raise PermissionError(13, "em uso")

    monkeypatch.setattr(xlsx_mod, "_replace", locked)
    summary, events, _ = engine_run(tmp_path, ["1"], Bio())
    assert summary.output.status is OutputStatus.FAILED and summary.output.detail == "locked"  # type: ignore[attr-defined]
    assert summary.status is RunStatus.COMPLETED_WITH_FAILURES  # type: ignore[attr-defined]
    assert "OUTPUT_FAILED" in events.codes()
