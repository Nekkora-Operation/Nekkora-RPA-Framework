"""Escrita GENÉRICA de planilha ``.xlsx`` (0.3.0). Requer o extra ``excel`` (openpyxl).

Contrato (``nekkora-rpa/xlsx-output`` 1.0.0):

- **Esquema declarado**: cabeçalhos e ordem das colunas (``OutputColumn``), campo-chave para duplicidade, nome da aba,
  limite de linhas. Cada coluna tem tipo (``text``, ``iso_date``, ``url``), obrigatoriedade e limite de caracteres.
- **Células sempre TEXTO** (tipo ``s`` + formato ``@``): um valor que começa com ``=``, ``+``, ``-``, ``@`` nunca vira
  fórmula (prevenção de *formula injection*) e ainda recebe o prefixo de citação do Excel. Saneamento: Unicode NFC,
  espaços colapsados, caracteres de controle proibidos no XML removidos.
- **Gravação atômica**: arquivo temporário na mesma pasta → conferência REABRINDO o temporário → ``os.replace``.
  Arquivo travado (aberto no Excel), falha de E/S ou conferência divergente → o arquivo final anterior fica intacto e
  o temporário é apagado. Política de sobrescrita explícita (``OverwritePolicy``).
- **Destino transacional** (``XlsxResultSink``): ``deliver`` prepara a linha de cada transação bem-sucedida num estado
  persistido (JSON atômico na pasta da execução); ``finalize`` grava a planilha uma única vez. Com
  ``CommitPolicy.ALL_OR_NOTHING`` (padrão) nada é gravado se alguma transação não teve sucesso ou se o lote não chegou
  ao fim — nunca há planilha parcial. A retomada reaproveita as linhas já preparadas (idempotente por
  ``transaction_id``; duplicidade pela chave recusada) e reconcilia uma consolidação interrompida pelo hash.
- **Privacidade**: nenhum valor de célula vai para eventos ou mensagens de erro; os erros são tokens
  (``XlsxOutputError.code``). O estado e a planilha SÃO dados de negócio (como os resultados JSON) e ficam só nos
  caminhos indicados por quem configura o processo.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
from collections.abc import Callable, Mapping, Sequence
from dataclasses import dataclass
from datetime import date
from enum import StrEnum
from pathlib import Path
from typing import Any

from .core.checkpoints import atomic_write
from .core.contracts import (
    OutputReport,
    OutputRun,
    OutputStatus,
    SinkRejected,
    TransactionItem,
    TransactionOutcome,
    TransactionStatus,
)
from .observability.safety import is_token
from .text import normalize_text

__all__ = [
    "STATE_SCHEMA",
    "STATE_VERSION",
    "ColumnKind",
    "CommitPolicy",
    "OutputColumn",
    "OverwritePolicy",
    "XlsxOutputError",
    "XlsxResultSink",
    "XlsxSchema",
    "prepare_row",
    "read_back",
    "write_xlsx",
]

STATE_SCHEMA = "nekkora-rpa/xlsx-output"
STATE_VERSION = "1.0.0"
MAX_ROWS = 100_000
MAX_COLUMNS = 100
MAX_CELL = 32_000  # o Excel aceita 32.767 caracteres por célula
_HEADER_MAX = 100
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_URL = re.compile(r"https?://[^\s]{1,1990}")
_ISO = re.compile(r"\d{4}-\d{2}-\d{2}")
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")
_replace = os.replace  # troca atômica do arquivo FINAL (isolada para os testes de trava e de queda)


class XlsxOutputError(Exception):
    """Falha estruturada da escrita: ``code`` é um token (``schema_invalid``, ``locked``, ``row_invalid``...)."""

    def __init__(self, code: str) -> None:
        self.code = code if is_token(code) else "invalid"
        super().__init__(self.code)


class ColumnKind(StrEnum):
    TEXT = "text"
    ISO_DATE = "iso_date"
    """Texto ``AAAA-MM-DD`` de uma data existente (gravado como texto: datas anteriores a 1900 não existem no Excel)."""
    URL = "url"


class OverwritePolicy(StrEnum):
    REPLACE = "replace"
    """Substitui o arquivo existente (atômico)."""
    FAIL_IF_EXISTS = "fail_if_exists"
    """Recusa se já existe um arquivo que não foi esta mesma consolidação que gravou."""


class CommitPolicy(StrEnum):
    ALL_OR_NOTHING = "all_or_nothing"
    """Grava só se TODAS as transações tiveram sucesso e o lote chegou ao fim."""
    SUCCEEDED_ONLY = "succeeded_only"
    """Grava as linhas bem-sucedidas mesmo com falhas (a planilha deixa de representar a entrada inteira)."""


@dataclass(frozen=True)
class OutputColumn:
    header: str
    """Texto do cabeçalho (como aparece na planilha)."""
    field: str
    """Nome do campo na saída da transação (``ProcessResult.output``)."""
    kind: ColumnKind = ColumnKind.TEXT
    required: bool = True
    max_chars: int = 2000


@dataclass(frozen=True)
class XlsxSchema:
    columns: tuple[OutputColumn, ...]
    key_field: str
    """Campo que identifica a linha (duplicidade recusada)."""
    sheet: str = "Dados"
    max_rows: int = 10_000

    def __post_init__(self) -> None:
        cols = tuple(self.columns)
        object.__setattr__(self, "columns", cols)
        if not 1 <= len(cols) <= MAX_COLUMNS:
            raise XlsxOutputError("schema_invalid")
        headers: set[str] = set()
        fields: set[str] = set()
        for c in cols:
            if not isinstance(c, OutputColumn) or not isinstance(c.header, str) or not isinstance(c.field, str):
                raise XlsxOutputError("schema_invalid")
            h = normalize_text(c.header)
            if (
                not h
                or h != c.header
                or len(h) > _HEADER_MAX
                or h.startswith(_FORMULA_START)
                or _ILLEGAL.search(h)
                or h.casefold() in headers
            ):
                raise XlsxOutputError("schema_header_invalid")
            headers.add(h.casefold())
            if not re.fullmatch(r"[a-z][a-z0-9_]{0,59}", c.field) or c.field in fields:
                raise XlsxOutputError("schema_field_invalid")
            fields.add(c.field)
            if not isinstance(c.kind, ColumnKind):
                raise XlsxOutputError("schema_kind_invalid")
            if isinstance(c.max_chars, bool) or not isinstance(c.max_chars, int) or not 1 <= c.max_chars <= MAX_CELL:
                raise XlsxOutputError("schema_limit_invalid")
        key = next((c for c in cols if c.field == self.key_field), None)
        if key is None or not key.required:
            raise XlsxOutputError("schema_key_invalid")
        if (
            not isinstance(self.sheet, str)
            or not 1 <= len(self.sheet) <= 31
            or _SHEET_BAD.search(self.sheet)
            or self.sheet != self.sheet.strip()
            or self.sheet.startswith("'")
        ):
            raise XlsxOutputError("schema_sheet_invalid")
        if isinstance(self.max_rows, bool) or not isinstance(self.max_rows, int) or not 1 <= self.max_rows <= MAX_ROWS:
            raise XlsxOutputError("schema_limit_invalid")

    @property
    def headers(self) -> tuple[str, ...]:
        return tuple(c.header for c in self.columns)

    def revision(self) -> str:
        doc = {
            "sheet": self.sheet,
            "key": self.key_field,
            "max_rows": self.max_rows,
            "columns": [[c.header, c.field, c.kind.value, c.required, c.max_chars] for c in self.columns],
        }
        canonical = json.dumps(doc, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        return "sha256:" + hashlib.sha256(canonical.encode("utf-8")).hexdigest()


# --------------------------------------------------------------------------------------------- linhas
def _clean(value: object, column: OutputColumn) -> str:
    if value is None:
        text = ""
    elif isinstance(value, str):
        text = value
    elif isinstance(value, date) and column.kind is ColumnKind.ISO_DATE:
        text = value.isoformat()
    else:
        raise XlsxOutputError(f"value_type_{column.field}")
    text = normalize_text(_ILLEGAL.sub("", text))
    if not text:
        if column.required:
            raise XlsxOutputError(f"missing_{column.field}")
        return ""
    if len(text) > column.max_chars:
        raise XlsxOutputError(f"too_long_{column.field}")
    if column.kind is ColumnKind.ISO_DATE:
        if not _ISO.fullmatch(text):
            raise XlsxOutputError(f"date_invalid_{column.field}")
        try:
            date.fromisoformat(text)
        except ValueError:
            raise XlsxOutputError(f"date_invalid_{column.field}") from None
    elif column.kind is ColumnKind.URL and not _URL.fullmatch(text):
        raise XlsxOutputError(f"url_invalid_{column.field}")
    return text


def prepare_row(schema: XlsxSchema, data: Mapping[str, object]) -> tuple[str, ...]:
    """Valida e saneia UMA linha (na ordem das colunas). Campo fora do esquema → ``unknown_field``."""
    if not isinstance(data, Mapping):
        raise XlsxOutputError("row_invalid")
    known = {c.field for c in schema.columns}
    for k in data:
        if k not in known:
            raise XlsxOutputError("unknown_field")
    return tuple(_clean(data.get(c.field), c) for c in schema.columns)


def _check_rows(schema: XlsxSchema, rows: Sequence[tuple[str, ...]]) -> None:
    if len(rows) > schema.max_rows:
        raise XlsxOutputError("too_many_rows")
    idx = [c.field for c in schema.columns].index(schema.key_field)
    seen: set[str] = set()
    for r in rows:
        if r[idx] in seen:
            raise XlsxOutputError("duplicate_key")
        seen.add(r[idx])


# --------------------------------------------------------------------------------------------- arquivo
def _openpyxl() -> Any:
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError:
        raise XlsxOutputError("excel_extra_not_installed") from None
    return openpyxl


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            h.update(chunk)
    return "sha256:" + h.hexdigest()


def read_back(path: Path | str, schema: XlsxSchema) -> list[tuple[str, ...]]:
    """Reabre a planilha e devolve as linhas de dados como texto, CONFERINDO: uma única aba com o nome do esquema,
    cabeçalho idêntico, nenhuma fórmula, nenhuma célula fora das colunas. Divergência → ``XlsxOutputError``."""
    openpyxl = _openpyxl()
    try:
        wb = openpyxl.load_workbook(Path(path), read_only=True, data_only=False, keep_links=False)
    except Exception:  # noqa: BLE001 — arquivo ilegível/corrompido: só o código sai
        raise XlsxOutputError("unreadable") from None
    try:
        if list(wb.sheetnames) != [schema.sheet]:
            raise XlsxOutputError("verify_sheet")
        ws = wb[schema.sheet]
        out: list[tuple[str, ...]] = []
        width = len(schema.columns)
        for n, row in enumerate(ws.iter_rows()):
            cells = list(row)
            values: list[str] = []
            for i, cell in enumerate(cells):
                v = cell.value
                if getattr(cell, "data_type", "s") == "f" or (
                    isinstance(v, str) and v.startswith("=") and cell.data_type != "s"
                ):
                    raise XlsxOutputError("verify_formula")
                if i >= width:
                    if v not in (None, ""):
                        raise XlsxOutputError("verify_extra_cell")
                    continue
                if v is not None and not isinstance(v, str):
                    raise XlsxOutputError("verify_cell_type")
                values.append(v or "")
            values += [""] * (width - len(values))
            if n == 0:
                if tuple(values) != schema.headers:
                    raise XlsxOutputError("verify_header")
                continue
            out.append(tuple(values))
        return out
    finally:
        wb.close()


def _write_file(tmp: Path, schema: XlsxSchema, rows: Sequence[tuple[str, ...]]) -> None:
    openpyxl = _openpyxl()
    from openpyxl.styles import Font  # type: ignore[import-untyped]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = schema.sheet
    bold = Font(bold=True)
    for col, header in enumerate(schema.headers, start=1):
        c = ws.cell(row=1, column=col, value=header)
        c.data_type = "s"
        c.font = bold
    for r, values in enumerate(rows, start=2):
        for col, value in enumerate(values, start=1):
            if value == "":
                continue
            c = ws.cell(row=r, column=col)
            c.value = value
            c.data_type = "s"  # SEMPRE texto: "=..." não vira fórmula
            c.number_format = "@"
            if value.startswith(_FORMULA_START):
                c.quotePrefix = True  # o Excel mostra e edita como texto
    ws.freeze_panes = "A2"
    with tmp.open("wb") as fh:
        wb.save(fh)
        fh.flush()
        os.fsync(fh.fileno())


def _commit(
    path: Path,
    schema: XlsxSchema,
    rows: Sequence[tuple[str, ...]],
    overwrite: OverwritePolicy,
    own_sha: str | None,
    before_replace: Callable[[str], None] | None = None,
) -> str:
    """Grava ``rows`` em ``path`` de forma atômica e conferida; devolve o sha256 do arquivo gravado.
    ``before_replace(sha)`` roda depois da conferência e ANTES da troca (o destino registra o hash pretendido)."""
    _check_rows(schema, rows)
    if path.suffix.lower() != ".xlsx":
        raise XlsxOutputError("path_not_xlsx")
    if not path.parent.is_dir():
        raise XlsxOutputError("folder_missing")
    if path.exists():
        if path.is_dir():
            raise XlsxOutputError("path_is_folder")
        if overwrite is OverwritePolicy.FAIL_IF_EXISTS and (own_sha is None or _sha256(path) != own_sha):
            raise XlsxOutputError("exists")
    tmp = path.with_name(f".{path.stem}.{os.getpid()}.tmp.xlsx")
    try:
        try:
            _write_file(tmp, schema, rows)
        except XlsxOutputError:
            raise
        except OSError:
            raise XlsxOutputError("io_error") from None
        if read_back(tmp, schema) != [tuple(r) for r in rows]:
            raise XlsxOutputError("verify_mismatch")
        sha = _sha256(tmp)
        if before_replace is not None:
            before_replace(sha)
        try:
            _replace(tmp, path)
        except PermissionError:
            raise XlsxOutputError("locked") from None  # Windows: arquivo aberto no Excel
        except OSError:
            raise XlsxOutputError("io_error") from None
        return sha
    finally:
        try:
            if tmp.exists():
                tmp.unlink()
        except OSError:
            pass


def write_xlsx(
    path: Path | str,
    schema: XlsxSchema,
    rows: Sequence[Mapping[str, object]],
    *,
    overwrite: OverwritePolicy = OverwritePolicy.FAIL_IF_EXISTS,
) -> OutputReport:
    """Escrita direta (sem transações): valida TODAS as linhas antes de tocar no disco e grava de forma atômica."""
    prepared = [prepare_row(schema, r) for r in rows]
    _commit(Path(path), schema, prepared, overwrite, None)
    return OutputReport(OutputStatus.COMMITTED, len(prepared))


# --------------------------------------------------------------------------------------------- destino transacional
class XlsxResultSink:
    """``ResultSink`` com consolidação no fim (``FinalizingResultSink``) que grava UMA planilha.

    - ``deliver``: transação bem-sucedida → linha validada e preparada no estado (``SinkRejected`` com token se a
      linha é inválida, incompleta ou repete a chave de outra transação; nada é preparado). Transação sem sucesso →
      registrada como impedimento (``ALL_OR_NOTHING``).
    - ``finalize``: grava a planilha (atômica, conferida) ou a RETÉM com o motivo; nunca propaga exceção.
    - ``expected_rows``: quantidade exata de linhas exigida para gravar (ex.: o tamanho da lista de entrada).
    """

    idempotent = True

    def __init__(
        self,
        path: Path | str,
        schema: XlsxSchema,
        *,
        state_dir: Path | str,
        overwrite: OverwritePolicy = OverwritePolicy.REPLACE,
        commit: CommitPolicy = CommitPolicy.ALL_OR_NOTHING,
        expected_rows: int | None = None,
    ) -> None:
        if not isinstance(schema, XlsxSchema):
            raise XlsxOutputError("schema_invalid")
        if expected_rows is not None and (
            isinstance(expected_rows, bool)
            or not isinstance(expected_rows, int)
            or not 1 <= expected_rows <= schema.max_rows
        ):
            raise XlsxOutputError("expected_rows_invalid")
        self.path = Path(path)
        if self.path.suffix.lower() != ".xlsx":
            raise XlsxOutputError("path_not_xlsx")
        self.schema = schema
        self.overwrite = OverwritePolicy(overwrite)
        self.commit_policy = CommitPolicy(commit)
        self.expected_rows = expected_rows
        self.state_path = Path(state_dir) / f"{self.path.stem}.xlsx-output.json"
        self._rows: dict[str, list[str]] = {}
        self._blocked: dict[str, str] = {}
        self._state = "OPEN"
        self._committed_sha: str | None = None
        self._pending_sha: str | None = None
        self._load()

    # ---------------------------------------------------------------- estado persistido
    @property
    def state(self) -> str:
        """``OPEN``, ``COMMITTING``, ``COMMITTED``, ``WITHHELD`` ou ``FAILED`` (observável e persistido)."""
        return self._state

    @property
    def staged_rows(self) -> int:
        return len(self._rows)

    def _load(self) -> None:
        if not self.state_path.exists():
            return
        try:
            doc = json.loads(self.state_path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            raise XlsxOutputError("state_unreadable") from None
        if (
            not isinstance(doc, dict)
            or doc.get("schema") != STATE_SCHEMA
            or doc.get("version") != STATE_VERSION
            or doc.get("schema_revision") != self.schema.revision()
            or doc.get("target") != self.path.name
        ):
            raise XlsxOutputError("state_mismatch")  # outra planilha/esquema: nunca mistura linhas
        rows = doc.get("rows")
        blocked = doc.get("blocked")
        if not isinstance(rows, dict) or not isinstance(blocked, dict):
            raise XlsxOutputError("state_unreadable")
        width = len(self.schema.columns)
        for tx, values in rows.items():
            if not is_token(tx) or not isinstance(values, list) or len(values) != width:
                raise XlsxOutputError("state_unreadable")
            self._rows[tx] = [str(v) for v in values]
        self._blocked = {str(k): str(v) for k, v in blocked.items()}
        self._state = str(doc.get("state", "OPEN"))
        self._committed_sha = doc.get("committed_sha256")
        self._pending_sha = doc.get("pending_sha256")
        if self._state == "COMMITTING":
            # consolidação interrompida entre a conferência e o fim: o hash decide se o arquivo final é o pretendido
            if self._pending_sha and self.path.exists() and _sha256(self.path) == self._pending_sha:
                self._state, self._committed_sha = "COMMITTED", self._pending_sha
            else:
                self._state = "OPEN"
            self._pending_sha = None
            self._save()

    def _save(self) -> None:
        doc = {
            "schema": STATE_SCHEMA,
            "version": STATE_VERSION,
            "schema_revision": self.schema.revision(),
            "target": self.path.name,
            "state": self._state,
            "rows": self._rows,
            "blocked": self._blocked,
            "committed_sha256": self._committed_sha,
            "pending_sha256": self._pending_sha,
        }
        self.state_path.parent.mkdir(parents=True, exist_ok=True)
        atomic_write(self.state_path, json.dumps(doc, ensure_ascii=False, indent=1) + "\n")

    # ---------------------------------------------------------------- ResultSink
    def deliver(self, item: TransactionItem, outcome: TransactionOutcome) -> None:
        tx = outcome.transaction_id
        if outcome.status is not TransactionStatus.SUCCEEDED:
            if tx in self._rows:
                raise SinkRejected("already_staged")
            self._blocked[tx] = outcome.status.value
            self._save()
            return
        try:
            values = list(prepare_row(self.schema, outcome.output or {}))
        except XlsxOutputError as err:
            raise SinkRejected(err.code) from None
        idx = [c.field for c in self.schema.columns].index(self.schema.key_field)
        for other, vals in self._rows.items():
            if other != tx and vals[idx] == values[idx]:
                raise SinkRejected("duplicate_key")
        if tx not in self._rows and len(self._rows) >= self.schema.max_rows:
            raise SinkRejected("too_many_rows")
        changed = self._rows.get(tx) != values
        self._rows[tx] = values
        self._blocked.pop(tx, None)
        if changed and self._state != "OPEN":
            self._state = "OPEN"
        self._save()

    def finalize(self, run: OutputRun) -> OutputReport:
        try:
            return self._finalize(run)
        except XlsxOutputError as err:
            return self._fail(err.code)
        except Exception:  # noqa: BLE001 — a consolidação nunca derruba o engine; só o motivo genérico sai
            return self._fail("unexpected")

    def _fail(self, code: str) -> OutputReport:
        self._state = "FAILED"
        self._pending_sha = None
        try:
            self._save()
        except OSError:
            pass
        return OutputReport(OutputStatus.FAILED, len(self._rows), code)

    def _withhold(self, code: str) -> OutputReport:
        self._state = "WITHHELD"
        self._save()
        return OutputReport(OutputStatus.WITHHELD, len(self._rows), code)

    def _finalize(self, run: OutputRun) -> OutputReport:
        for tx, status in run.statuses:
            if status is not TransactionStatus.SUCCEEDED and tx not in self._rows:
                self._blocked.setdefault(tx, status.value)
        if self.commit_policy is CommitPolicy.ALL_OR_NOTHING:
            if not run.complete:
                return self._withhold("incomplete_run")
            if self._blocked:
                return self._withhold("transactions_not_succeeded")
        if not self._rows:
            return self._withhold("no_rows")
        if self.expected_rows is not None and len(self._rows) != self.expected_rows:
            return self._withhold("row_count_mismatch")
        rows = [tuple(v) for v in self._rows.values()]
        if (
            self._state == "COMMITTED"
            and self._committed_sha
            and self.path.exists()
            and _sha256(self.path) == self._committed_sha
        ):
            return OutputReport(OutputStatus.COMMITTED, len(rows), "already_committed")

        def pending(sha: str) -> None:
            self._state, self._pending_sha = "COMMITTING", sha
            self._save()

        sha = _commit(self.path, self.schema, rows, self.overwrite, self._committed_sha, pending)
        self._state, self._committed_sha, self._pending_sha = "COMMITTED", sha, None
        self._save()
        return OutputReport(OutputStatus.COMMITTED, len(rows))
